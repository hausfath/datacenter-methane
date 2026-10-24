"""Grid reference intensities: eGRID2023 rev2 BA averages with symmetric upstream terms, new grid NGCC, Cambium 2024 LRMER.

Writes data/grid_inputs.json (small derived file, enough to regenerate everything without the xlsx files).
"""
import json
from pathlib import Path

import numpy as np
import openpyxl

import emissions as em

ROOT = Path(__file__).resolve().parents[1]
LB = 0.45359237
REGIONS = {'ERCOT': 'ERCO', 'PJM': 'PJM'}          # eGRID BA codes; US from US23
CAMBIUM_REGIONS = {'ERCOT': ['ERCOT'], 'PJM': ['PJM_East', 'PJM_West']}


def _sheet(wb, name, key):
    rows = list(wb[name].iter_rows(values_only=True))
    hdr = rows[1]
    return {r[hdr.index(key)]: dict(zip(hdr, r)) for r in rows[2:] if r[hdr.index(key)]}


def egrid_inputs():
    wb = openpyxl.load_workbook(ROOT / 'data/egrid/egrid2023_data_rev2.xlsx', read_only=True, data_only=True)
    ba = _sheet(wb, 'BA23', 'BACODE')
    us = list(_sheet(wb, 'US23', 'YEAR').values())[0]
    srl = _sheet(wb, 'SRL23', 'SUBRGN')
    out = {}

    def one(d, p, name):
        ghr = d[p + 'GCO2RT'] / d[p + 'GCO2R']                        # gas MMBtu per gas MWh
        chr_ = d[p + 'CCO2RT'] / d[p + 'CCO2R'] if d[p + 'CCO2R'] else 0.0
        out[name] = dict(
            co2e_kg_mwh=d[p + 'C2ERTA'] * LB, co2_kg_mwh=d[p + 'CO2RTA'] * LB,
            nonbaseload_co2e_kg_mwh=d[p + 'NBC2E'] * LB,
            gas_share=d[p + 'GSPR'], coal_share=d[p + 'CLPR'],
            gas_mmbtu_per_mwh=d[p + 'GSPR'] * ghr, coal_mmbtu_per_mwh=d[p + 'CLPR'] * chr_,
            gas_hr=ghr, coal_hr=chr_)
    for name, code in REGIONS.items():
        one(ba[code], 'BA', name)
    one(us, 'US', 'US')
    for s in ['RFCE', 'RFCW', 'ERCT']:
        one(srl[s], 'SR', f'sub_{s}')
    return out


def cambium_lrmer(start=2027, years=20):
    """Simple-mean (0% discount) LRMER, combustion + precombustion, AR6 GWP100, end-use basis, kg CO2e/MWh."""
    wb = openpyxl.load_workbook(ROOT / 'data/cambium/Cambium24_Workbook.xlsx', read_only=True, data_only=True)
    rows = list(wb['Data - Annual'].iter_rows(values_only=True))
    scen_row, year_row = rows[3], rows[4]
    data = {r[1]: r for r in rows[5:] if r[1]}
    blocks = dict(co2c=15, ch4c=63, n2oc=111, co2p=159, ch4p=207, n2op=255)
    scen = [scen_row[15 + 6 * k] for k in range(8)]
    yrs = np.array(year_row[15:21], dtype=float)
    target = np.arange(start, start + years)
    g = em.GWP['AR6_100']
    out = {}
    for reg, gea_list in CAMBIUM_REGIONS.items():
        out[reg] = {}
        for k, s in enumerate(scen):
            vals = []
            for gea in gea_list:
                r = data[gea]
                ts = {b: np.interp(target, yrs, np.array(r[c + 6 * k:c + 6 * k + 6], dtype=float)) for b, c in blocks.items()}
                co2e = ts['co2c'] + ts['co2p'] + (ts['ch4c'] + ts['ch4p']) / 1000 * g['ch4'] + (ts['n2oc'] + ts['n2op']) / 1000 * g['n2o']
                vals.append(co2e.mean())
            out[reg][s] = float(np.mean(vals))   # PJM: unweighted mean of East and West (no load weights in workbook)
    return out


if __name__ == '__main__':
    res = dict(egrid=egrid_inputs(), cambium_lrmer_2027_2046=cambium_lrmer(),
               notes='eGRID2023 rev2 (6/12/2025). Gas/coal MMBtu per grid MWh = share x (CO2 output rate / CO2 input rate). '
                     'Cambium 2024 (Aug 2024 policies, pre-OBBBA), 2027-2046 simple mean, combustion+precombustion, AR6 GWP100, end-use.')
    json.dump(res, open(ROOT / 'data/grid_inputs.json', 'w'), indent=1)
    for k, v in res['egrid'].items():
        print(k, {a: round(b, 3) for a, b in v.items()})
    for k, v in res['cambium_lrmer_2027_2046'].items():
        print(k, {a: round(b, 1) for a, b in v.items()})
